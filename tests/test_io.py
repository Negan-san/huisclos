"""Vérifie les lectures et exports sur des documents entièrement fictifs.

Les fichiers TXT, DOCX et PDF sont fabriqués dans le dossier temporaire du test.
La fixture réseau bloque les connexions durant tous les allers-retours.
"""

from dataclasses import FrozenInstanceError
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter

from huisclos.io.readers import DocumentReadError, read_document, read_pasted_text
from huisclos.io.writers import to_docx_bytes, to_txt_bytes
from huisclos.models import Detection, EntityType, Source
from huisclos.pseudonymize.mapping import Mapping
from huisclos.pseudonymize.reidentify import EXACT, restore
from huisclos.pseudonymize.tokens import pseudonymize

pytestmark = pytest.mark.usefixtures("no_network")


def _read_file(path: Path) -> bytes:
    """Renvoie les octets du fichier temporaire reçu."""
    return path.read_bytes()


def _pdf_file(tmp_path: Path, texts: tuple[str, ...]) -> Path:
    """Crée un PDF temporaire avec une page par texte, ou un rectangle si vide."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in texts:
        pdf.add_page()
        if text:
            pdf.cell(w=180, h=10, text=text)
        else:
            pdf.rect(20, 20, 40, 30)
    path = tmp_path / "fiction.pdf"
    pdf.output(path)
    return path


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1252"])
def test_txt_encodings(tmp_path: Path, encoding: str) -> None:
    """Retrouve les accents et l'apostrophe du TXT, avec avertissement si cp1252."""
    text = "Le dossier fictif est rédigé pour l\u2019étude."
    path = tmp_path / "fiction.TXT"
    path.write_bytes(text.encode(encoding))
    result = read_document(path.name, _read_file(path))
    assert result.text == text
    assert result.source_type == "txt"
    assert result.char_count == len(text)
    assert result.warnings == (("encoding_cp1252",) if encoding == "cp1252" else ())


@pytest.mark.parametrize("source", ["txt", "paste", "docx"])
def test_normalization_and_nonbreaking_spaces(tmp_path: Path, source: str) -> None:
    """Normalise NFC et CR/LF tout en conservant les espaces insécables du texte."""
    text = "  Etude re\u0301dige\u0301e\r\nArticle\u00a012\rMontant\u202f: 20\n"
    expected = "  Etude rédigée\nArticle\u00a012\nMontant\u202f: 20\n"
    if source == "paste":
        result = read_pasted_text(text)
    else:
        path = tmp_path / f"fiction.{source}"
        data = text.encode("utf-8") if source == "txt" else to_docx_bytes(text, "")
        path.write_bytes(data)
        result = read_document(path.name, _read_file(path))
        if source == "docx":
            expected = "\n" + expected
    assert result.text == expected
    assert result.char_count == len(expected)


def test_read_result_is_frozen() -> None:
    """Vérifie que le résultat collé ne peut pas être modifié après validation."""
    result = read_pasted_text("Texte fictif")
    assert result.source_type == "paste"
    assert result.warnings == ()
    with pytest.raises(FrozenInstanceError):
        result.text = "Autre texte"


def test_docx_body_table_headers_and_footers(tmp_path: Path) -> None:
    """Lit corps et tableau, puis chaque en-tête et pied distinct une seule fois."""
    document = Document()
    document.add_paragraph("Paragraphe fictif")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Alpha", "Beta"), ("Gamma", "Delta")), strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    section = document.sections[0]
    section.header.paragraphs[0].text = "Cabinet fictif, 12 rue des Nuages, 75001 Paris"
    section.header.paragraphs[0].text += " - 01 00 00 00 00"
    section.footer.paragraphs[0].text = "Pied de page fictif"
    section.first_page_header.paragraphs[0].text = "Première page fictive"
    section.even_page_header.paragraphs[0].text = "Page paire fictive"
    section.first_page_footer.paragraphs[0].text = "Pied première page"
    section.even_page_footer.paragraphs[0].text = "Pied page paire"
    document.add_section()
    last_section = document.add_section()
    last_section.header.is_linked_to_previous = False
    last_section.header.paragraphs[0].text = "Autre en-tête fictif"
    path = tmp_path / "fiction.docx"
    document.save(path)
    result = read_document(path.name, _read_file(path))
    assert result.source_type == "docx"
    assert result.warnings == ()
    assert result.text.splitlines() == [
        "Paragraphe fictif",
        "Alpha | Beta",
        "Gamma | Delta",
        "",
        "",
        "Cabinet fictif, 12 rue des Nuages, 75001 Paris - 01 00 00 00 00",
        "Première page fictive",
        "Page paire fictive",
        "Pied de page fictif",
        "Pied première page",
        "Pied page paire",
        "Autre en-tête fictif",
    ]
    assert result.char_count == len(result.text)


def test_docx_merged_cells_and_header_table(tmp_path: Path) -> None:
    """Ne répète pas une cellule fusionnée et inclut les tableaux d'en-tête."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "Cellule fusionnée fictive"
    header = document.sections[0].header
    header_table = header.add_table(rows=1, cols=1, width=1000000)
    header_table.cell(0, 0).text = "Adresse fictive dans un tableau"
    path = tmp_path / "fusion.docx"
    document.save(path)
    result = read_document(path.name, _read_file(path))
    assert result.text.count("Cellule fusionnée fictive") == 1
    assert result.text.count("Adresse fictive dans un tableau") == 1


def test_pdf_two_pages(tmp_path: Path) -> None:
    """Retrouve les textes natifs de deux pages séparés par une ligne vide."""
    texts = (
        "Premiere page du dossier juridique fictif.",
        "Deuxieme page du dossier juridique fictif.",
    )
    path = _pdf_file(tmp_path, texts)
    result = read_document(path.name, _read_file(path))
    assert result.text == "\n\n".join(texts)
    assert result.source_type == "pdf"
    assert result.warnings == ()
    assert result.char_count == len(result.text)


@pytest.mark.parametrize("texts", [("",), (" " * 30,)])
def test_pdf_without_enough_useful_text(tmp_path: Path, texts: tuple[str, ...]) -> None:
    """Refuse uniquement une extraction entièrement dépourvue de texte utile."""
    path = _pdf_file(tmp_path, texts)
    with pytest.raises(DocumentReadError) as error:
        read_document(path.name, _read_file(path))
    assert error.value.code == "pdf_no_text"


def test_pdf_useful_text_boundary(tmp_path: Path) -> None:
    """Accepte le texte disponible mais signale la page vide malgré une moyenne suffisante."""
    path = _pdf_file(tmp_path, ("a" * 40, ""))
    result = read_document(path.name, _read_file(path))
    assert result.text == "a" * 40 + "\n\n"
    assert result.warnings == ("pdf_empty_pages:2",)


@pytest.mark.parametrize("password", ["secret-fictif", ""])
def test_encrypted_pdf(tmp_path: Path, password: str) -> None:
    """Refuse aussi le PDF chiffré dont le mot de passe utilisateur est vide."""
    path = _pdf_file(tmp_path, ("Document juridique entierement fictif.",))
    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(path))
    writer.encrypt(user_password=password, owner_password="proprietaire-fictif")
    encrypted_path = tmp_path / "chiffre.pdf"
    writer.write(encrypted_path)
    with pytest.raises(DocumentReadError) as error:
        read_document(encrypted_path.name, _read_file(encrypted_path))
    assert error.value.code == "pdf_encrypted"


@pytest.mark.parametrize(
    ("filename", "data", "code"),
    [
        ("fiction.rtf", b"texte", "unsupported_type"),
        ("sans_extension", b"texte", "unsupported_type"),
        ("fiction.txt", b"x" * 6_000_000, "too_large"),
        ("fiction.txt", b"x" * 200_001, "too_long"),
        ("fiction.txt", b"", "empty"),
        ("fiction.docx", b"", "empty"),
        ("fiction.pdf", b"", "empty"),
        ("fiction.txt", b" \r\n\t", "empty"),
        ("fiction.txt", b"\xef\xbb\xbf", "empty"),
        ("fiction.txt", b"\x81", "encoding"),
        ("fiction.docx", b"\x13\xa7\x00\x9c", "docx_invalid"),
        ("fiction.pdf", b"\x13\xa7\x00\x9c", "pdf_invalid"),
    ],
    ids=[
        "unsupported-rtf",
        "unsupported-no-extension",
        "too-large",
        "too-long",
        "empty-txt",
        "empty-docx",
        "empty-pdf",
        "empty-whitespace",
        "empty-bom",
        "encoding",
        "docx-invalid",
        "pdf-invalid",
    ],
)
def test_document_errors(tmp_path: Path, filename: str, data: bytes, code: str) -> None:
    """Expose les codes prévus pour les formats, encodages et limites invalides."""
    path = tmp_path / filename
    path.write_bytes(data)
    with pytest.raises(DocumentReadError) as error:
        read_document(path.name, _read_file(path))
    assert error.value.code == code
    assert str(error.value) == code


@pytest.mark.parametrize(
    "text", ["", " \r\n\u00a0", "x" * 200_001], ids=["empty", "whitespace", "too-long"]
)
def test_pasted_text_errors(text: str) -> None:
    """Refuse le texte collé vide ou au-delà de la limite commune."""
    with pytest.raises(DocumentReadError) as error:
        read_pasted_text(text)
    assert error.value.code == ("too_long" if len(text) > 200_000 else "empty")


def test_limits_include_boundary_and_normalization(tmp_path: Path) -> None:
    """Accepte les bornes inclusives et mesure les caractères après normalisation."""
    path = tmp_path / "limite.txt"
    path.write_bytes(b"x" * 200_000)
    assert read_document(path.name, _read_file(path)).char_count == 200_000
    assert read_pasted_text("x" * 200_000).char_count == 200_000
    assert read_pasted_text("e\u0301" * 200_000).char_count == 200_000
    assert read_document("f.txt", b"abc", max_bytes=3, max_chars=3).text == "abc"
    with pytest.raises(DocumentReadError) as error:
        read_document("f.txt", b"abcd", max_bytes=3)
    assert error.value.code == "too_large"
    with pytest.raises(DocumentReadError) as error:
        read_document("f.txt", b"abcd", max_chars=3)
    assert error.value.code == "too_long"


@pytest.mark.parametrize("source", ["docx", "pdf"])
def test_extracted_text_too_long(tmp_path: Path, source: str) -> None:
    """Applique aussi la limite personnalisée au texte extrait de DOCX et PDF."""
    if source == "pdf":
        path = _pdf_file(tmp_path, ("Document juridique entierement fictif.",))
    else:
        path = tmp_path / "long.docx"
        path.write_bytes(to_docx_bytes("Texte fictif", "Titre"))
    with pytest.raises(DocumentReadError) as error:
        read_document(path.name, _read_file(path), max_chars=5)
    assert error.value.code == "too_long"


def test_empty_docx(tmp_path: Path) -> None:
    """Refuse un DOCX valide qui ne contient aucun texte."""
    path = tmp_path / "vide.docx"
    Document().save(path)
    with pytest.raises(DocumentReadError) as error:
        read_document(path.name, _read_file(path))
    assert error.value.code == "empty"


def test_txt_roundtrip(tmp_path: Path) -> None:
    """Conserve exactement le texte UTF-8 et ses lignes, sans ajouter de BOM."""
    text = "Texte fictif rédigé\n\nAdresse\u00a0: 12 rue des Nuages\n"
    path = tmp_path / "export.txt"
    data = to_txt_bytes(text)
    path.write_bytes(data)
    assert data == text.encode("utf-8")
    assert not data.startswith(b"\xef\xbb\xbf")
    assert read_document(path.name, _read_file(path)).text == text


@pytest.mark.parametrize("text", ["Ligne fictive\n\nDeuxième ligne\n", "", "a\r\nb\rc"])
def test_docx_roundtrip_and_metadata(tmp_path: Path, text: str) -> None:
    """Conserve titre et lignes dans le DOCX et vide les métadonnées personnelles."""
    title = "Dossier fictif"
    path = tmp_path / "export.docx"
    data = to_docx_bytes(text, title)
    path.write_bytes(data)
    document = Document(BytesIO(data))
    expected_lines = [title, *text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    assert [paragraph.text for paragraph in document.paragraphs] == expected_lines
    assert read_document(path.name, _read_file(path)).text == "\n".join(expected_lines)
    properties = document.core_properties
    assert properties.author == ""
    assert properties.last_modified_by == ""
    assert properties.comments == ""
    assert properties.title == ""


@pytest.mark.parametrize("source", ["docx", "pdf"])
def test_generated_document_pseudonymization_json_and_restore(tmp_path: Path, source: str) -> None:
    """Relie la lecture d'un fichier fabriqué à la restauration depuis son JSON."""
    values = ("Jean DUPONT", "Marie MARTIN", "Paul DURAND")
    if source == "pdf":
        path = _pdf_file(tmp_path, ("Dossier fictif : " + ", ".join(values),))
    else:
        document = Document()
        document.add_paragraph("Dossier fictif : " + values[0])
        document.add_table(rows=1, cols=1).cell(0, 0).text = values[1]
        document.sections[0].header.paragraphs[0].text = values[2]
        path = tmp_path / "controle.docx"
        document.save(path)
    text = read_document(path.name, path.read_bytes()).text
    detections = []
    for value in values:
        assert text.count(value) == 1
        start = text.index(value)
        detections.append(
            Detection(start, start + len(value), EntityType.PERSON, 1.0, Source.MANUAL, value)
        )
    result = pseudonymize(text, detections)
    assert all(value not in result.text for value in values)
    restored = restore(result.text, Mapping.from_json(result.mapping.to_json()))
    assert restored.mode == EXACT
    assert restored.text == text
