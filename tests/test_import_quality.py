"""Vérifie la complétude des imports avec des documents fictifs en mémoire.

Les PDF et DOCX synthétiques couvrent ordre, pages douteuses et structures omises.
Aucun binaire ni accès externe n'est nécessaire avant la revue humaine.
"""

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import parse_xml
from fpdf import FPDF
from PIL import Image

from huisclos.io.read_quality import ReadQuality, warning_parameters
from huisclos.io.readers import DocumentReadError, ReadResult, read_document, read_pasted_text

pytestmark = pytest.mark.usefixtures("no_network")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def pdf_bytes(texts: tuple[str | None, ...]) -> bytes:
    """Fabrique une page par texte ; None représente une page image sans texte natif."""
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in texts:
        pdf.add_page()
        if text is None:
            pdf.image(Image.new("RGB", (32, 32), "black"), x=10, y=10, w=30)
        elif text:
            pdf.cell(w=180, h=10, text=text)
    return bytes(pdf.output())


def docx_bytes(structure: str = "", parts: dict[str, str] | None = None) -> bytes:
    """Fabrique un DOCX lisible avec un bloc XML et des parties internes facultatives."""
    document = Document()
    document.add_paragraph("Visible fictional body")
    if structure:
        document._element.body.insert(1, parse_xml(f'<w:p xmlns:w="{W}">{structure}</w:p>'))
    buffer = BytesIO()
    document.save(buffer)
    if not parts:
        return buffer.getvalue()
    output = BytesIO()
    with ZipFile(buffer) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in original.infolist():
            target.writestr(entry, original.read(entry.filename))
        for name, xml in parts.items():
            target.writestr(name, xml)
    return output.getvalue()


def test_mixed_pdf_cannot_hide_image_page_in_average() -> None:
    """Signale exactement la page image entre deux pages riches sans inventer son contenu."""
    text = "Fictional contract " * 10
    result = read_document("mixed.pdf", pdf_bytes((text, None, text)))
    assert result.quality == ReadQuality(empty_pages=(2,))
    assert result.warnings == ("pdf_empty_pages:2",)
    assert result.text == f"{text}\n\n\n\n{text}"
    assert result.char_count == len(result.text)


@pytest.mark.parametrize("texts,pages", [(("Signed",), (1,)), (("a" * 20, "End"), (2,))])
def test_short_pdf_page_is_accepted(texts: tuple[str, ...], pages: tuple[int, ...]) -> None:
    """Une page courte légitime reste disponible pour la revue, avec état inconnu."""
    result = read_document("short.pdf", pdf_bytes(texts))
    assert result.text == "\n\n".join(texts)
    assert result.quality.status == "unknown"
    assert result.quality.short_pages == pages
    assert result.quality.empty_pages == ()
    assert result.warnings == (f"pdf_short_pages:{','.join(map(str, pages))}",)


def test_pdf_page_numbers_include_blank_and_short_pages() -> None:
    """Sépare les pages sans texte des pages courtes en conservant les indices du fichier."""
    result = read_document("pages.pdf", pdf_bytes(("", "End", "x" * 20, None, "OK")))
    assert result.quality.empty_pages == (1, 4)
    assert result.quality.short_pages == (2, 5)
    assert result.warnings == ("pdf_empty_pages:1,4", "pdf_short_pages:2,5")


@pytest.mark.parametrize("texts", [(None,), ("", " "), (None, None)])
def test_pdf_without_extracted_text_keeps_error(texts: tuple[str | None, ...]) -> None:
    """Refuse les extractions vides sans attribuer une cause ni exposer de contenu."""
    with pytest.raises(DocumentReadError, match="^pdf_no_text$") as error:
        read_document("empty.pdf", pdf_bytes(texts))
    assert error.value.code == "pdf_no_text"
    assert error.value.quality.empty_pages == tuple(range(1, len(texts) + 1))


def test_pdf_native_text_does_not_assert_completeness() -> None:
    """Un texte natif assez long garde les anciens champs sans garantir tout le PDF."""
    text = "Fictional selectable legal text."
    result = read_document("text.pdf", pdf_bytes((text,)))
    assert (result.text, result.source_type, result.warnings, result.char_count) == (
        text,
        "pdf",
        (),
        len(text),
    )
    assert result.quality.status == "unknown"


def test_docx_order_nested_tables_and_headers() -> None:
    """Conserve l'ordre des blocs du corps, des cellules imbriquées et de l'en-tête."""
    document = Document()
    document.add_paragraph("Before")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Cell before"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    cell.add_paragraph("Cell after")
    document.add_paragraph("After")
    header = document.sections[0].header
    header.paragraphs[0].text = "Header before"
    header.add_table(rows=1, cols=1, width=1000000).cell(0, 0).text = "Header table"
    header.add_paragraph("Header after")
    document.sections[0].footer.paragraphs[0].text = "Footer"
    buffer = BytesIO()
    document.save(buffer)
    result = read_document("order.docx", buffer.getvalue())
    assert result.text.splitlines() == [
        "Before",
        "Cell before",
        "Nested",
        "",
        "Cell after",
        "After",
        "Header before",
        "Header table",
        "Header after",
        "Footer",
    ]
    assert result.warnings == ()
    assert result.quality.status == "unknown"


@pytest.mark.parametrize(
    "structure,code",
    [
        ('<w:r><w:footnoteReference w:id="1"/></w:r>', "notes"),
        ('<w:r><w:endnoteReference w:id="1"/></w:r>', "notes"),
        ('<w:commentRangeStart w:id="0"/>', "comments"),
        (
            "<w:r><w:pict><w:txbxContent><w:p><w:r><w:t>Hidden fiction</w:t>"
            "</w:r></w:p></w:txbxContent></w:pict></w:r>",
            "text_boxes",
        ),
        ('<w:ins w:id="0"><w:r><w:t>Hidden fiction</w:t></w:r></w:ins>', "revisions"),
        ('<w:del w:id="0"><w:r><w:delText>Hidden fiction</w:delText></w:r></w:del>', "revisions"),
        ('<w:pPr><w:pPrChange w:id="0"/></w:pPr>', "revisions"),
        ('<w:r><w:cellIns w:id="0"/></w:r>', "revisions"),
        ('<w:r><w:numberingChange w:id="0"/></w:r>', "revisions"),
        ('<w:customXmlDelRangeStart w:id="0"/>', "revisions"),
        ("<w:r><w:object/></w:r>", "embedded_content"),
        ("<w:sdt><w:sdtContent/></w:sdt>", "embedded_content"),
    ],
)
def test_docx_unsupported_structure_is_explicit(structure: str, code: str) -> None:
    """Signale chaque catégorie sans prétendre lire le texte que Word y conserve."""
    result = read_document("limited.docx", docx_bytes(structure))
    assert result.quality.status == "partial"
    assert result.quality.unsupported_structures == (code,)
    assert result.warnings == (f"docx_{code}",)
    assert "Hidden fiction" not in result.text
    assert "Hidden fiction" not in repr(result.quality)


def test_docx_notes_comments_and_header_box_parts() -> None:
    """Repère notes, commentaires et zones de texte dans les parties XML internes."""
    parts = {
        "word/footnotes.xml": f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="1"/></w:footnotes>',
        "word/comments.xml": f'<w:comments xmlns:w="{W}"><w:comment w:id="0"/></w:comments>',
        "word/header9.xml": f'<w:hdr xmlns:w="{W}"><w:txbxContent/></w:hdr>',
    }
    result = read_document("parts.docx", docx_bytes(parts=parts))
    assert result.warnings == ("docx_comments", "docx_notes", "docx_text_boxes")


def test_docx_separator_is_not_an_omitted_note() -> None:
    """Ne signale pas les séparateurs de notes sans note ni renvoi effectif."""
    xml = f'<w:footnotes xmlns:w="{W}"><w:footnote w:type="separator" w:id="-1"/></w:footnotes>'
    result = read_document("separator.docx", docx_bytes(parts={"word/footnotes.xml": xml}))
    assert result.warnings == ()


def test_external_docx_attachment_is_not_opened() -> None:
    """Signale une inclusion externe sans suivre sa relation, sous garde réseau."""
    relation = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="external" TargetMode="External" Target="https://example.org/fiction" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk"/>'
        "</Relationships>"
    )
    document = Document()
    document.add_paragraph("Visible fiction")
    document._element.body.insert(
        1,
        parse_xml(
            f'<w:altChunk xmlns:w="{W}" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'r:id="external"/>'
        ),
    )
    buffer = BytesIO()
    document.save(buffer)
    output = BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            if entry.filename != "word/_rels/document.xml.rels":
                target.writestr(entry, source.read(entry.filename))
        target.writestr("word/_rels/document.xml.rels", relation)
    result = read_document("linked.docx", output.getvalue())
    assert result.text == "Visible fiction"
    assert result.warnings == ("docx_embedded_content",)


def test_read_result_and_text_compatibility() -> None:
    """Les quatre arguments historiques restent utilisables ; le texte simple est complet."""
    assert ReadResult("abc", "txt", (), 3).quality.status == "unknown"
    assert read_pasted_text("abc").quality.status == "complete"
    assert read_document("abc.txt", b"abc").quality.status == "complete"
    assert warning_parameters("encoding_cp1252") == ("encoding_cp1252", {})
    assert warning_parameters("pdf_empty_pages:2,4") == ("pdf_empty_pages", {"pages": "2, 4"})


@pytest.mark.parametrize("suffix", ["0", "-1", "<b>fiction</b>", "2,", "two", "\u0661"])
def test_page_warning_parameters_reject_content(suffix: str) -> None:
    """N'accepte que des entiers ASCII positifs dans les paramètres affichés."""
    warning = f"pdf_empty_pages:{suffix}"
    assert warning_parameters(warning) == (warning, {})


def test_docx_only_unsupported_text_keeps_warning_on_error() -> None:
    """Un Word dont tout le texte est omis reste refusé avec sa limite explicite."""
    document = Document()
    document._element.body.insert(
        0,
        parse_xml(
            f'<w:p xmlns:w="{W}"><w:ins w:id="0"><w:r><w:t>Hidden fiction</w:t></w:r></w:ins></w:p>'
        ),
    )
    buffer = BytesIO()
    document.save(buffer)
    with pytest.raises(DocumentReadError, match="^empty$") as error:
        read_document("revisions.docx", buffer.getvalue())
    assert error.value.warnings == ("docx_revisions",)
    assert error.value.quality.status == "partial"


def test_english_dev_build_is_pure_reproducible_and_compatible(tmp_path: Path) -> None:
    """Reconstruit uniquement le dev anglais et le valide avec le chargeur public existant."""
    from huisclos.evaluation.corpus import load_corpus
    from scripts.build_english_dev import ROOT, build_dataset

    sources = {
        path.name: path.read_text(encoding="utf-8") for path in (ROOT / "source").glob("*.md")
    }
    jsonl, manifest_text = build_dataset(sources)
    assert build_dataset(dict(reversed(list(sources.items())))) == (jsonl, manifest_text)
    assert (ROOT / "dev.jsonl").read_bytes() == jsonl.encode("utf-8")
    assert (ROOT / "manifest.json").read_bytes() == manifest_text.encode("utf-8")
    path = tmp_path / "english.jsonl"
    path.write_text(jsonl, encoding="utf-8")
    assert len(load_corpus(path)) == 8
    manifest = json.loads(manifest_text)
    assert manifest["dataset_id"] == "english_dev_v1"
    assert manifest["language"] == "en"
    assert manifest["split"] == "dev"
    assert all(250 <= doc["word_count"] <= 450 for doc in manifest["documents"])


@pytest.mark.parametrize(
    "body",
    [
        "[[Eden Frost|FOREIGN_ID]]",
        "[[|PERSON]]",
        "[[ Eden Frost|PERSON]]",
        "[[Eden [[Frost|PERSON]]|PERSON]]",
        "[[Eden Frost|PERSON]",
        "[[Eden Frost|PERSON]] and Eden Frost",
        "[[eden@other.invalid|EMAIL]]",
    ],
)
def test_english_dev_builder_rejects_invalid_annotations(body: str) -> None:
    """Refuse les balises, types, répétitions et courriels qui violent les conventions."""
    from scripts.build_english_dev import parse_source

    source = (
        "---\nid: en_dev_01\ndoc_type: case_memorandum\nsplit: dev\ntitle: Fiction\n---\n" + body
    )
    with pytest.raises(ValueError):
        parse_source(source)


def test_english_dev_builder_multiline_unicode_and_known_repeats() -> None:
    """Calcule les offsets Unicode et autorise toutes les occurrences annotées."""
    from scripts.build_english_dev import parse_source

    source = (
        "---\nid: en_dev_01\ndoc_type: case_memorandum\nsplit: dev\ntitle: Fiction\n---\n"
        "Ms [[Zoé Frost|PERSON]], [[12 rue des Ormes\n75013 Paris|ADDRESS]]. "
        "[[Zoé Frost|PERSON]] signed."
    )
    document = parse_source(source)
    for entity in document["entities"]:
        assert document["text"][entity["start"] : entity["end"]] == entity["text"]
    assert len(document["entities"]) == 3


def test_english_dev_builder_requires_eight_sources() -> None:
    """Le constructeur isolé refuse un lot manquant sans consulter un autre corpus."""
    from scripts.build_english_dev import build_dataset

    with pytest.raises(ValueError, match="Huit sources"):
        build_dataset({})
