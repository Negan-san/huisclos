"""Exporte localement le texte pseudonymisé ou restauré en fin de pipeline.

Le texte fourni devient des octets TXT ou un document DOCX simple.
Le DOCX conserve les lignes et ne renseigne aucune métadonnée personnelle.
"""

from io import BytesIO

from docx import Document


def to_txt_bytes(text: str) -> bytes:
    """Encode le texte fourni en UTF-8 sans ajouter de marque BOM."""
    return text.encode("utf-8")


def to_docx_bytes(text: str, title: str) -> bytes:
    """Renvoie un DOCX avec le titre fourni puis un paragraphe par ligne du texte."""
    document = Document()
    document.add_paragraph(title)
    normalized_lines = text.replace("\r\n", "\n").replace("\r", "\n")
    for line in normalized_lines.split("\n"):
        document.add_paragraph(line)
    # Le modèle fourni par la bibliothèque contient un auteur et un commentaire.
    properties = document.core_properties
    for name in (
        "author",
        "last_modified_by",
        "comments",
        "title",
        "subject",
        "keywords",
        "category",
        "identifier",
        "language",
        "content_status",
        "version",
    ):
        setattr(properties, name, "")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
